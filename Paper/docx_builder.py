"""Small helper layer over python-docx for the manuscript.

Inline markup in text strings:
  **bold**   *italic*   ~sub~   ^sup^   [@key] or [@key1;key2] citations
  {{HL:text}} yellow-highlighted text for items the authors must confirm
"""
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor, Inches

FONT = "Times New Roman"
TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*|~.+?~|\^.+?\^|\[@[^\]]+\]|\{\{HL:.+?\}\})")


class Manuscript:
    def __init__(self, refs, endnote=False):
        self.refs = refs              # key -> dict(text=..., url=..., endnote=...)
        self.endnote = endnote        # True: write EndNote temporary citations, no typed bibliography
        self.order = []               # citation order of first appearance
        self.fig_n = 0
        self.tab_n = 0
        self.eq_n = 0
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Cm(2.3))
        sec.top_margin = sec.bottom_margin = Cm(2.3)
        st = self.doc.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(11)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        pf = st.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.15
        for lvl, size in ((1, 13), (2, 11.5), (3, 11)):
            h = self.doc.styles[f"Heading {lvl}"]
            h.font.name = FONT
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.italic = lvl == 3
            h.font.color.rgb = RGBColor(0, 0, 0)
            h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            h.element.rPr.rFonts.set(qn("w:ascii"), FONT)
            h.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
            h.paragraph_format.space_before = Pt(12 if lvl == 1 else 8)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
        self._page_numbers()
        # python-docx's default template omits the required w:percent on w:zoom
        zoom = self.doc.settings.element.find(qn("w:zoom"))
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")

    # ---------------------------------------------------------------- utils
    def _page_numbers(self):
        p = self.doc.sections[0].footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            run._r.append(el)
        run.font.size = Pt(9)

    def cite(self, keys):
        nums = []
        for k in keys:
            k = k.strip()
            if k not in self.refs:
                raise KeyError(f"unknown citation key: {k}")
            if k not in self.order:
                self.order.append(k)
            nums.append(self.order.index(k) + 1)
        if self.endnote:
            # EndNote temporary citation, e.g. {Jukanti, 2016; Sharma, 2013}; EndNote's
            # "Update Citations and Bibliography" turns it into a real EndNote field.
            ordered = sorted(dict.fromkeys(k.strip() for k in keys), key=lambda k: self.order.index(k))
            return "{" + "; ".join(self.refs[k]["endnote"] for k in ordered) + "}"
        nums = sorted(set(nums))
        # compress runs of three or more consecutive numbers: 3,4,5 -> 3–5
        parts, i = [], 0
        while i < len(nums):
            j = i
            while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
                j += 1
            if j - i >= 2:
                parts.append(f"{nums[i]}–{nums[j]}")
            else:
                parts.extend(str(n) for n in nums[i:j + 1])
            i = j + 1
        return "[" + ", ".join(parts) + "]"

    def runs(self, p, text, size=None, bold=False, italic=False):
        for piece in TOKEN.split(text):
            if not piece:
                continue
            b, it, sub, sup, hl = bold, italic, False, False, False
            if piece.startswith("**"):
                piece, b = piece[2:-2], True
            elif piece.startswith("{{HL:"):
                piece, hl = piece[5:-2], True
            elif piece.startswith("[@"):
                piece = self.cite(piece[2:-1].split(";"))
            elif piece.startswith("*"):
                piece, it = piece[1:-1], True
            elif piece.startswith("~"):
                piece, sub = piece[1:-1], True
            elif piece.startswith("^"):
                piece, sup = piece[1:-1], True
            r = p.add_run(piece)
            r.bold, r.italic = b, it
            r.font.subscript, r.font.superscript = sub, sup
            if hl:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            if size:
                r.font.size = Pt(size)
        return p

    # -------------------------------------------------------------- blocks
    def h(self, text, level=1):
        self.doc.add_heading(text, level=level)

    def p(self, text, align="justify", size=None, italic=False, indent=False, after=None):
        para = self.doc.add_paragraph()
        para.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                          "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
        if indent:
            para.paragraph_format.first_line_indent = Cm(0.6)
        if after is not None:
            para.paragraph_format.space_after = Pt(after)
        return self.runs(para, text, size=size, italic=italic)

    def bullets(self, items, numbered=False):
        style = "List Number" if numbered else "List Bullet"
        for it in items:
            para = self.doc.add_paragraph(style=style)
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            para.paragraph_format.space_after = Pt(3)
            self.runs(para, it)

    def eq(self, text):
        self.eq_n += 1
        t = self.doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        w = [Cm(14.4), Cm(2.0)]
        for i, c in enumerate(t.rows[0].cells):
            c.width = w[i]
        left = t.rows[0].cells[0].paragraphs[0]
        left.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.runs(left, text)
        right = t.rows[0].cells[1].paragraphs[0]
        right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        right.add_run(f"({self.eq_n})")
        return self.eq_n

    def fig(self, path, caption, width_cm=16.0):
        self.fig_n += 1
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True
        para.add_run().add_picture(str(path), width=Cm(width_cm))
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        cap.paragraph_format.space_after = Pt(10)
        r = cap.add_run(f"Fig. {self.fig_n}. ")
        r.bold = True
        r.font.size = Pt(9.5)
        self.runs(cap, caption, size=9.5)
        return self.fig_n

    def table(self, caption, header, rows, widths_cm=None, font=8.5, note=None):
        self.tab_n += 1
        cap = self.doc.add_paragraph()
        cap.paragraph_format.keep_with_next = True
        cap.paragraph_format.space_before = Pt(6)
        r = cap.add_run(f"Table {self.tab_n}. ")
        r.bold = True
        r.font.size = Pt(9.5)
        self.runs(cap, caption, size=9.5)
        n = len(header)
        t = self.doc.add_table(rows=1 + len(rows), cols=n)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        if widths_cm is None:
            widths_cm = [16.4 / n] * n
        for ri, row in enumerate([header] + rows):
            for ci, val in enumerate(row):
                cell = t.rows[ri].cells[ci]
                cell.width = Cm(widths_cm[ci])
                para = cell.paragraphs[0]
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.line_spacing = 1.0
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
                self.runs(para, str(val), size=font, bold=(ri == 0))
        self._three_line(t)
        if note:
            np_ = self.doc.add_paragraph()
            np_.paragraph_format.space_after = Pt(10)
            self.runs(np_, note, size=8.5, italic=False)
        else:
            self.doc.add_paragraph().paragraph_format.space_after = Pt(4)
        return self.tab_n

    @staticmethod
    def _three_line(t):
        """Academic table rules: top, below-header, bottom."""
        def border(cell, edge, sz):
            tcPr = cell._tc.get_or_add_tcPr()
            b = tcPr.find(qn("w:tcBorders"))
            if b is None:
                b = OxmlElement("w:tcBorders")
                tcPr.append(b)
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:color"), "000000")
            b.append(e)
        last = len(t.rows) - 1
        for ci in range(len(t.columns)):
            border(t.rows[0].cells[ci], "top", 10)
            border(t.rows[0].cells[ci], "bottom", 6)
            border(t.rows[last].cells[ci], "bottom", 10)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def references(self):
        self.h("References", 1)
        if self.endnote:
            return  # EndNote inserts the formatted bibliography at the end of the document
        for i, k in enumerate(self.order, 1):
            ref = self.refs[k]
            para = self.doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.left_indent = Cm(0.8)
            para.paragraph_format.first_line_indent = Cm(-0.8)
            para.paragraph_format.space_after = Pt(3)
            # the link is rendered once, as a hyperlink, so drop any URL written into the text
            text = re.sub(r"\s*https?://\S+$", "", ref["text"]).rstrip()
            self.runs(para, f"[{i}]\t" + text + " ", size=9.5)
            if ref.get("url"):
                self._hyperlink(para, ref["url"])

    def _hyperlink(self, para, url):
        part = para.part
        r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                              is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), r_id)
        r = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        c = OxmlElement("w:color")
        c.set(qn("w:val"), "1F4E9A")
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "19")
        rPr.extend([c, u, sz])
        r.append(rPr)
        t = OxmlElement("w:t")
        t.text = url
        r.append(t)
        h.append(r)
        para._p.append(h)

    def save(self, path):
        unused = [k for k in self.refs if k not in self.order]
        self.doc.save(str(path))
        return unused


def read_md_table(path, keep=None, rename=None, transform=None):
    """Parse a GitHub-markdown table produced by build_tables.py."""
    lines = [l.strip() for l in open(path, encoding="utf-8") if l.strip().startswith("|")]
    rows = [[c.strip() for c in l.strip("|").split("|")] for l in lines]
    header, body = rows[0], [r for r in rows[2:]]
    idx = list(range(len(header))) if keep is None else [header.index(k) for k in keep]
    header = [header[i] for i in idx]
    body = [[r[i] for i in idx] for r in body]
    if rename:
        header = [rename.get(h, h) for h in header]
    if transform:
        body = [transform(r) for r in body]
    return header, body
